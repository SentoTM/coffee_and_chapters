"""Stars Hollow Book Club (repo coffee_and_chapters) — "Tinder" de libros para elegir el reto de lectura.

Flask + SQLite. Dos usuarios definidos en la variable APP_USERS.

Cada usuario tiene, por libro:
  - decision: want (leer) · reject (descartar) · skip (pasar, vuelve a salir) · NULL (sin decidir)
  - is_read:  casilla aparte "ya lo he leído"
"""
import csv
import io
import json
import os
import random
import sqlite3
from datetime import date
from functools import wraps
from urllib.parse import quote_plus

from flask import (Flask, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("DATABASE_PATH", os.path.join(BASE_DIR, "data", "app.db"))
SEED_FILE = os.environ.get("SEED_FILE", os.path.join(BASE_DIR, "seed", "coffee_and_chapters.xlsx"))
BLURBS_FILE = os.path.join(BASE_DIR, "seed", "sinopsis.tsv")  # Nº <TAB> sinopsis breve
DECISIONS = ("want", "reject", "skip")
# Libros que no entran nunca: nivel 7 (descartables) y duplicados/solapados por Nº
# (112 Inferno ya está dentro de 235 La Divina Comedia).
EXCLUDED_LEVEL_PREFIX = "7"
EXCLUDED_IDS = {x.strip() for x in os.environ.get("EXCLUDED_IDS", "112").split(",") if x.strip()}
SCHEMA_VERSION = 5

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-cambia-esto")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
app.json.sort_keys = False  # mantener el orden de los campos de la tarjeta


# --------------------------------------------------------------------------- #
# Usuarios: APP_USERS="sento:clave1,and:clave2"
# --------------------------------------------------------------------------- #
def load_users():
    raw = os.environ.get("APP_USERS", "sento:sento,and:and")
    users = {}
    for pair in raw.split(","):
        if ":" in pair:
            name, pwd = pair.split(":", 1)
            users[name.strip().lower()] = generate_password_hash(pwd.strip())
    return users


USERS = load_users()
# Quién ve las herramientas de administración: el buscador en Descubrir y la página Importar
SEARCH_USERS = {u.strip().lower() for u in os.environ.get("SEARCH_USERS", "sento").split(",") if u.strip()}

# Columnas del Excel con el estado previo de cada persona → usuario de la app
# IMPORT_USER_COLUMNS="Vicen:sento,Andrea:and"
USER_COLUMNS = {
    k.strip().lower(): v.strip().lower()
    for k, v in (p.split(":", 1) for p in os.environ.get("IMPORT_USER_COLUMNS", "Vicen:sento,Andrea:and").split(",") if ":" in p)
}


def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if session.get("user") not in USERS:
            if request.path.startswith("/api/"):
                return jsonify(error="no autenticado"), 401
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapper


@app.context_processor
def inject_permissions():
    return {"can_search": session.get("user") in SEARCH_USERS}


def others(user):
    """El resto de usuarios, en el orden de APP_USERS."""
    return [u for u in USERS if u != user]


N_USERS = len(USERS)
USER_PH = ",".join("?" * len(USERS))  # placeholders para "username IN (...)"
# nº de usuarios actuales que quieren leer cada libro
WANTS_SQL = f"(SELECT COUNT(*) FROM votes w WHERE w.book_id = b.id AND w.decision = 'want' AND w.username IN ({USER_PH}))"


# --------------------------------------------------------------------------- #
# Base de datos
# --------------------------------------------------------------------------- #
SCHEMA = """
CREATE TABLE IF NOT EXISTS books (
    id       INTEGER PRIMARY KEY AUTOINCREMENT,
    ext_id   TEXT UNIQUE,                 -- Nº del Excel, si existe
    title    TEXT NOT NULL,               -- título a mostrar (español si hay)
    original TEXT NOT NULL DEFAULT '',    -- título original
    author   TEXT NOT NULL DEFAULT '',
    level    TEXT NOT NULL DEFAULT '',
    extra    TEXT NOT NULL DEFAULT '{}',
    blurb    TEXT NOT NULL DEFAULT '',     -- sinopsis breve (seed/sinopsis.tsv)
    added    TEXT NOT NULL DEFAULT (datetime('now')),
    UNIQUE (original, author)
);
CREATE TABLE IF NOT EXISTS votes (
    username TEXT NOT NULL,
    book_id  INTEGER NOT NULL REFERENCES books(id) ON DELETE CASCADE,
    decision TEXT CHECK (decision IN ('want','reject','skip')),
    is_read  INTEGER NOT NULL DEFAULT 0,
    updated  TEXT NOT NULL DEFAULT (datetime('now')),
    PRIMARY KEY (username, book_id)
);
-- Lecturas del reto: libros elegidos de entre los de "Nos lo quedamos"
CREATE TABLE IF NOT EXISTS picks (
    book_id   INTEGER PRIMARY KEY REFERENCES books(id) ON DELETE CASCADE,
    picked_by TEXT NOT NULL,
    status    TEXT NOT NULL DEFAULT 'next' CHECK (status IN ('next','reading','done')),
    turn      INTEGER NOT NULL DEFAULT 1,  -- 0 = apuntada directamente como leída: no cuenta para el turno
    added     TEXT NOT NULL DEFAULT (datetime('now')),
    updated   TEXT NOT NULL DEFAULT (datetime('now'))
);
"""


def connect():
    os.makedirs(os.path.dirname(DB_PATH) or ".", exist_ok=True)
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")
    return db


def get_db():
    if "db" not in g:
        g.db = connect()
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = connect()
    version = db.execute("PRAGMA user_version").fetchone()[0]
    if version < 2:
        # Versión 1 (sin datos reales): se recrea.
        db.executescript("DROP TABLE IF EXISTS votes; DROP TABLE IF EXISTS books;")
    db.executescript(SCHEMA)
    if version in (2, 3):
        migrate_to_v4(db)
    db.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")
    db.commit()
    purge_excluded(db)
    empty = db.execute("SELECT COUNT(*) FROM books").fetchone()[0] == 0
    if empty and os.path.exists(SEED_FILE):
        with open(SEED_FILE, "rb") as f:
            res = import_books(db, os.path.basename(SEED_FILE), f.read())
        print(f"[seed] {res}")
    if version < 5:
        demo_vote(db)
    sync_blurbs(db)
    cols = {r["name"] for r in db.execute("PRAGMA table_info(picks)")}
    if "turn" not in cols:
        db.execute("ALTER TABLE picks ADD COLUMN turn INTEGER NOT NULL DEFAULT 1")
        db.commit()
    if "grp" not in cols:
        # Grupo del reto: 'all' (todos) o una pareja 'and+sento'. Lo existente queda en 'all'.
        db.execute("ALTER TABLE picks ADD COLUMN grp TEXT NOT NULL DEFAULT 'all'")
        db.commit()
    db.close()


def load_blurbs():
    if not os.path.exists(BLURBS_FILE):
        return {}
    out = {}
    with open(BLURBS_FILE, encoding="utf-8") as f:
        for line in f:
            n, _, text = line.rstrip("\n").partition("\t")
            if n.strip() and text.strip():
                out[n.strip()] = text.strip()
    return out


def sync_blurbs(db):
    """Copia las sinopsis de seed/sinopsis.tsv a la base en cada arranque (también en bases ya cargadas)."""
    cols = {r["name"] for r in db.execute("PRAGMA table_info(books)")}
    if "blurb" not in cols:
        db.execute("ALTER TABLE books ADD COLUMN blurb TEXT NOT NULL DEFAULT ''")
    db.executemany("UPDATE books SET blurb = ? WHERE ext_id = ?",
                   [(text, n) for n, text in load_blurbs().items()])
    db.commit()


def demo_vote(db):
    """Una sola vez: And marca 'leer' Ana Karenina (nº 22) para poder probar el match."""
    row = db.execute("SELECT id FROM books WHERE ext_id = '22'").fetchone()
    if not row or "and" not in USERS:
        return
    db.execute("INSERT OR IGNORE INTO votes (username, book_id) VALUES ('and', ?)", (row["id"],))
    db.execute("""UPDATE votes SET decision = 'want', updated = datetime('now')
                  WHERE username = 'and' AND book_id = ? AND decision IS NULL""", (row["id"],))
    db.commit()


def migrate_to_v4(db):
    """Versiones 2-3 precargaban decisiones desde el Excel. Ahora solo se precarga 'leído':
    se quitan las decisiones que siguen tal cual las dejó la carga inicial
    (lo que hayáis deslizado después se respeta)."""
    cur = db.execute(
        """UPDATE votes SET decision = NULL
           WHERE decision IS NOT NULL
             AND julianday(updated) - julianday((SELECT added FROM books b WHERE b.id = votes.book_id)) < 60.0/86400""")
    db.execute("DELETE FROM votes WHERE decision IS NULL AND is_read = 0")
    db.commit()
    print(f"[migrate v4] {cur.rowcount} decisiones precargadas deshechas")


def purge_excluded(db):
    """Quita de la base los libros excluidos (y sus votos) si ya estaban cargados."""
    ph = ",".join("?" * len(EXCLUDED_IDS)) or "NULL"
    db.execute(f"DELETE FROM books WHERE level LIKE ? OR ext_id IN ({ph})",
               [EXCLUDED_LEVEL_PREFIX + "%", *EXCLUDED_IDS])
    db.commit()


def is_excluded(ext_id, level):
    return (ext_id in EXCLUDED_IDS) or level.startswith(EXCLUDED_LEVEL_PREFIX)


# --------------------------------------------------------------------------- #
# Importación
# --------------------------------------------------------------------------- #
def _norm(s):
    return str(s or "").strip().lower()


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _sheet_rows(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    headers = [_clean(h) or f"col{i+1}" for i, h in enumerate(rows[0])]
    return [dict(zip(headers, r + [None] * (len(headers) - len(r)))) for r in rows[1:]]


def read_records(filename, data):
    """Lista de dicts {cabecera: valor}. Para el Excel del reto une 'Criba' + 'Lista' por Nº."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = {ws.title.lower(): ws for ws in wb.worksheets}
        if "criba" in sheets:
            main = _sheet_rows(sheets["criba"])
            if "lista" in sheets:
                by_n = {_clean(r.get("Nº")): r for r in _sheet_rows(sheets["lista"])}
                for r in main:
                    for k, v in by_n.get(_clean(r.get("Nº")), {}).items():
                        r.setdefault(k, v)
            return main
        return _sheet_rows(wb.active)
    if name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        rows = list(csv.reader(io.StringIO(text), dialect))
        headers = [h.strip() for h in rows[0]] if rows else []
        return [dict(zip(headers, r)) for r in rows[1:] if any(c.strip() for c in r)]
    raise ValueError("Formato no soportado: sube un .xlsx o .csv")


def _pick(rec, *keys):
    low = {_norm(k): k for k in rec}
    for k in keys:
        if k in low and _clean(rec[low[k]]):
            return _clean(rec[low[k]])
    return ""


# Campos que se muestran en la tarjeta, en este orden (el resto se ignora)
CARD_FIELDS = [
    ("Año", ("año", "año (1ª publ.)")),
    ("Tipo", ("tipo",)),
    ("Género", ("género", "genero")),
    ("Nivel", ("nivel", "nivel criba")),
    ("Estación", ("estación", "estacion")),
    ("País", ("país del autor", "país")),
    ("Idioma", ("idioma original",)),
    ("Premios", ("premios destacados", "premios")),
    ("Por qué", ("por qué está aquí",)),
]


def read_from_cell(value):
    """Columna de una persona en el Excel → ¿lo ha leído? (Leído / Releer = sí).
    El resto (Pendiente, Leyendo…) no se usa: las decisiones se toman deslizando."""
    v = _norm(value)
    return 1 if v.startswith(("leído", "leido", "releer")) else 0


def import_books(db, filename, data):
    records = read_records(filename, data)
    added = existing = excluded = 0
    for rec in records:
        original = _pick(rec, "obra", "titulo", "título", "title", "libro")
        if not original:
            continue
        spanish = _pick(rec, "título en español", "titulo en español")
        author = _pick(rec, "autor", "autora", "author", "autores")
        ext_id = _pick(rec, "nº", "n", "id") or None
        level = _pick(rec, "nivel", "nivel criba")
        if is_excluded(ext_id, level):
            excluded += 1
            continue
        extra = {label: _pick(rec, *keys) for label, keys in CARD_FIELDS}
        extra = {k: v for k, v in extra.items() if v}
        if not CARD_FIELDS or not extra:  # CSV genérico: muestra todo lo demás
            skip = {"obra", "titulo", "título", "title", "autor", "author", "nº"}
            extra = {k: _clean(v) for k, v in rec.items() if _clean(v) and _norm(k) not in skip}

        cur = db.execute(
            """INSERT OR IGNORE INTO books (ext_id, title, original, author, level, extra)
               VALUES (?,?,?,?,?,?)""",
            (ext_id, spanish or original, original, author, level, json.dumps(extra, ensure_ascii=False)),
        )
        if not cur.rowcount:
            existing += 1
            continue
        added += 1
        book_id = cur.lastrowid

        # Del Excel solo se toma quién lo ha leído (solo para libros nuevos).
        low = {_norm(k): k for k in rec}
        for user in USERS:
            col = next((c for c, u in USER_COLUMNS.items() if u == user), None)
            if col in low and read_from_cell(rec.get(low[col])):
                db.execute("INSERT OR IGNORE INTO votes (username, book_id, is_read) VALUES (?,?,1)",
                           (user, book_id))
    db.commit()
    return {"added": added, "existing": existing, "excluded": excluded}


# --------------------------------------------------------------------------- #
# Serialización
# --------------------------------------------------------------------------- #
def votes_for(db, book_ids):
    if not book_ids:
        return {}
    ph = ",".join("?" * len(book_ids))
    out = {}
    for r in db.execute(f"SELECT * FROM votes WHERE book_id IN ({ph})", list(book_ids)):
        out.setdefault(r["book_id"], {})[r["username"]] = {"decision": r["decision"], "read": bool(r["is_read"])}
    return out


HIDDEN_FIELDS = {"País", "Idioma", "Estación"}  # no se muestran como fila en la tarjeta


def book_dict(row, votes, user):
    extra = json.loads(row["extra"] or "{}")
    season = extra.get("Estación", "")
    q = f'{row["original"] or row["title"]} {row["author"]}'.strip()
    v = votes.get(row["id"], {})
    empty = {"decision": None, "read": False}
    return {
        "id": row["id"],
        "n": row["ext_id"],
        "title": row["title"],
        "original": row["original"] if row["original"] != row["title"] else "",
        "author": row["author"],
        "level": row["level"],
        "season": season,
        "blurb": row["blurb"] if "blurb" in row.keys() else "",
        "extra": {k: v for k, v in extra.items() if k not in HIDDEN_FIELDS},
        "goodreads": "https://www.goodreads.com/search?q=" + quote_plus(q),
        "me": v.get(user, empty),
        "others": [{"user": u, **v.get(u, empty)} for u in others(user)],
    }


def stats(db, user):
    total = db.execute("SELECT COUNT(*) FROM books").fetchone()[0]
    s = {"total": total, "want": 0, "reject": 0, "skip": 0}
    for r in db.execute(
            "SELECT decision, COUNT(*) c FROM votes WHERE username=? AND decision IS NOT NULL GROUP BY decision", (user,)):
        s[r["decision"]] = r["c"]
    s["pending"] = total - s["want"] - s["reject"] - s["skip"]
    s["read"] = db.execute("SELECT COUNT(*) FROM votes WHERE username=? AND is_read=1", (user,)).fetchone()[0]
    counts = db.execute(
        f"""SELECT SUM(n = ?) AS full, SUM(n >= 2 AND n < ?) AS almost
            FROM (SELECT {WANTS_SQL} AS n FROM books b)""", [N_USERS, N_USERS, *USERS]).fetchone()
    s["match"] = (counts["full"] or 0) if N_USERS >= 2 else 0
    s["almost"] = (counts["almost"] or 0) if N_USERS >= 3 else 0
    return s


# --------------------------------------------------------------------------- #
# Vistas
# --------------------------------------------------------------------------- #
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = _norm(request.form.get("user"))
        if user in USERS and check_password_hash(USERS[user], request.form.get("password", "")):
            session.permanent = True
            session["user"] = user
            nxt = request.args.get("next") or "/"
            return redirect(nxt if nxt.startswith("/") else "/")
        flash("Usuario o contraseña incorrectos")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def swipe():
    return render_template("swipe.html", user=session["user"])


@app.route("/listas")
@login_required
def lists():
    return render_template("lists.html", user=session["user"], others=others(session["user"]), many=N_USERS >= 3)


@app.route("/importar", methods=["GET", "POST"])
@login_required
def import_view():
    if session["user"] not in SEARCH_USERS:
        abort(403)
    db = get_db()
    if request.method == "POST":
        f = request.files.get("file")
        if not f or not f.filename:
            flash("Selecciona un fichero")
        else:
            try:
                r = import_books(db, f.filename, f.read())
                flash(f"Importados {r['added']} libros nuevos ({r['existing']} ya existían, {r['excluded']} excluidos)")
            except Exception as e:  # noqa: BLE001
                flash(f"Error al importar: {e}")
        return redirect(url_for("import_view"))
    total = db.execute("SELECT COUNT(*) FROM books").fetchone()[0]
    return render_template("import.html", user=session["user"], total=total)


# --------------------------------------------------------------------------- #
# API
# --------------------------------------------------------------------------- #
@app.route("/api/next")
@login_required
def api_next():
    """Siguiente libro: primero los no decididos (al azar); luego los pasados, del más antiguo."""
    user = session["user"]
    db = get_db()
    exclude = [int(x) for x in request.args.get("exclude", "").split(",") if x.isdigit()]
    not_in = f"AND b.id NOT IN ({','.join('?' * len(exclude))})" if exclude else ""

    ids = [r["id"] for r in db.execute(
        f"""SELECT b.id FROM books b
            LEFT JOIN votes v ON v.book_id = b.id AND v.username = ?
            WHERE v.decision IS NULL {not_in}""", [user, *exclude])]
    book_id = random.choice(ids) if ids else None
    if book_id is None:
        for excl, params in ((not_in, [user, *exclude]), ("", [user])):
            r = db.execute(
                f"""SELECT b.id FROM books b JOIN votes v ON v.book_id = b.id AND v.username = ?
                    WHERE v.decision = 'skip' {excl} ORDER BY v.updated LIMIT 1""", params).fetchone()
            if r:
                book_id = r["id"]
                break
    book = None
    if book_id is not None:
        row = db.execute("SELECT * FROM books WHERE id=?", (book_id,)).fetchone()
        book = book_dict(row, votes_for(db, [book_id]), user)
    return jsonify(book=book, stats=stats(db, user))


@app.route("/api/search")
@login_required
def api_search():
    """Busca por título, título original o autor (máx. 8)."""
    if session["user"] not in SEARCH_USERS:
        abort(403)
    q = request.args.get("q", "").strip()
    if len(q) < 2:
        return jsonify(books=[])
    like = f"%{q}%"
    rows = get_db().execute(
        """SELECT * FROM books WHERE title LIKE ? OR original LIKE ? OR author LIKE ?
           ORDER BY title LIMIT 8""", (like, like, like)).fetchall()
    votes = votes_for(get_db(), [r["id"] for r in rows])
    return jsonify(books=[book_dict(r, votes, session["user"]) for r in rows])


def upsert_vote(db, user, book_id, **fields):
    if not db.execute("SELECT 1 FROM books WHERE id=?", (book_id,)).fetchone():
        abort(404)
    db.execute("INSERT OR IGNORE INTO votes (username, book_id) VALUES (?,?)", (user, book_id))
    for k, v in fields.items():
        db.execute(f"UPDATE votes SET {k}=?, updated=datetime('now') WHERE username=? AND book_id=?", (v, user, book_id))
    db.commit()


@app.route("/api/vote", methods=["POST"])
@login_required
def api_vote():
    """{book_id, decision: want|reject|skip|null}  (null = deshacer / sin decidir)"""
    data = request.get_json(force=True, silent=True) or {}
    decision = data.get("decision")
    if decision is not None and decision not in DECISIONS:
        abort(400)
    db = get_db()
    user = session["user"]
    upsert_vote(db, user, data.get("book_id"), decision=decision)
    # match: "full" si lo queréis todos, "partial" si al menos 2 (con 3 o más usuarios)
    match, with_users, missing = None, [], []
    if decision == "want" and N_USERS >= 2:
        wanting = {r["username"] for r in db.execute(
            f"SELECT username FROM votes WHERE book_id=? AND decision='want' AND username IN ({USER_PH})",
            [data.get("book_id"), *USERS])}
        with_users = [u for u in others(user) if u in wanting]
        missing = [u for u in USERS if u not in wanting]
        if not missing:
            match = "full"
        elif with_users:
            match = "partial"
    return jsonify(ok=True, match=match, with_users=with_users, missing=missing, stats=stats(db, user))


@app.route("/api/read", methods=["POST"])
@login_required
def api_read():
    """{book_id, read: true|false}"""
    data = request.get_json(force=True, silent=True) or {}
    db = get_db()
    upsert_vote(db, session["user"], data.get("book_id"), is_read=1 if data.get("read") else 0)
    return jsonify(ok=True, stats=stats(db, session["user"]))


LIST_QUERIES = {
    # v1 = yo · WANTS_SQL = cuántos lo quieren leer
    "match": f"{WANTS_SQL} = {N_USERS}",
    "almost": f"{WANTS_SQL} >= 2 AND {WANTS_SQL} < {N_USERS}",
    "want": "v1.decision='want'",
    "skip": "v1.decision='skip'",
    "reject": "v1.decision='reject'",
    "read": "v1.is_read=1",
    "all": "1=1",
}


@app.route("/api/list/<kind>")
@login_required
def api_list(kind):
    if kind not in LIST_QUERIES:
        abort(404)
    user = session["user"]
    db = get_db()
    where = LIST_QUERIES[kind]
    params = [user] + list(USERS) * where.count("w.username IN")
    rows = db.execute(
        f"""SELECT b.* FROM books b
            LEFT JOIN votes v1 ON v1.book_id=b.id AND v1.username=?
            WHERE {where}
            ORDER BY b.level, b.title""",
        params,
    ).fetchall()
    votes = votes_for(db, [r["id"] for r in rows])
    return jsonify(books=[book_dict(r, votes, user) for r in rows])


# --------------------------------------------------------------------------- #
# Reto: elegir qué leer de entre los que os quedáis
# --------------------------------------------------------------------------- #
PICK_STATUS = ("next", "reading", "done")
# Los que queréis leer todos (params: USERS)
MATCH_SQL = f"SELECT b.* FROM books b WHERE {WANTS_SQL} = {N_USERS}"


# Grupos del reto: 'all' = todos; con 3+ usuarios, además una pareja por cada otro usuario ('and+sento')
def group_key(members):
    members = sorted(members)
    return "all" if len(members) == N_USERS else "+".join(members)


def group_members(key):
    return sorted(USERS) if key == "all" else sorted(key.split("+"))


def groups_for(user):
    out = [{"key": "all", "members": sorted(USERS)}]
    if N_USERS >= 3:
        out += [{"key": group_key([user, u]), "members": sorted([user, u]), "with": u} for u in others(user)]
    return out


def group_pool_sql(key):
    """Libros que quieren leer exactamente los miembros del grupo (params: los devuelve también)."""
    members = group_members(key)
    if key == "all":
        return MATCH_SQL, list(USERS)
    ph = ",".join("?" * len(members))
    sql = f"""SELECT b.* FROM books b
        WHERE (SELECT COUNT(*) FROM votes w WHERE w.book_id = b.id AND w.decision = 'want'
                 AND w.username IN ({ph})) = {len(members)}
          AND {WANTS_SQL} = {len(members)}"""
    return sql, [*members, *USERS]


def check_group(user, key):
    key = (key or "all").replace(" ", "+")  # un '+' sin codificar en la URL llega como espacio
    if key not in {g["key"] for g in groups_for(user)}:
        abort(403)
    return key


def current_season(today=None):
    m = (today or date.today()).month
    return {12: "Invierno", 1: "Invierno", 2: "Invierno", 3: "Primavera", 4: "Primavera", 5: "Primavera",
            6: "Verano", 7: "Verano", 8: "Verano"}.get(m, "Otoño")


def whose_turn(db, key="all"):
    """Elegís por turnos dentro de cada grupo, rotando en orden alfabético."""
    last = db.execute("SELECT picked_by FROM picks WHERE turn = 1 AND grp = ? ORDER BY added DESC, rowid DESC LIMIT 1",
                      (key,)).fetchone()
    users = group_members(key)
    if not users:
        return None
    if not last or last["picked_by"] not in users:
        return users[0]
    return users[(users.index(last["picked_by"]) + 1) % len(users)]


@app.route("/reto")
@login_required
def reto():
    return render_template("reto.html", user=session["user"], others=others(session["user"]),
                           groups=groups_for(session["user"]))


@app.route("/api/reto")
@login_required
def api_reto():
    user = session["user"]
    db = get_db()
    key = check_group(user, request.args.get("group"))
    sql, params = group_pool_sql(key)
    rows = db.execute(sql, params).fetchall()
    taken = {r["book_id"] for r in db.execute("SELECT book_id FROM picks WHERE grp != ?", (key,))}
    rows = [r for r in rows if r["id"] not in taken]  # ya elegido en otro grupo
    picks = {r["book_id"]: r for r in db.execute("SELECT * FROM picks WHERE grp = ?", (key,))}
    in_rows = {r["id"] for r in rows}
    extra_ids = [i for i in picks if i not in in_rows]
    if extra_ids:  # lecturas apuntadas a mano que no son coincidencias
        rows += db.execute(f"SELECT * FROM books WHERE id IN ({','.join('?' * len(extra_ids))})", extra_ids).fetchall()
    votes = votes_for(db, [r["id"] for r in rows])
    pool, chosen = [], []
    for r in rows:
        b = book_dict(r, votes, user)
        p = picks.get(r["id"])
        if p:
            b["pick"] = {"status": p["status"], "by": p["picked_by"], "added": p["added"], "updated": p["updated"],
                         "manual": not p["turn"]}
            chosen.append(b)
        else:
            pool.append(b)
    chosen.sort(key=lambda b: (PICK_STATUS.index(b["pick"]["status"]) if b["pick"]["status"] != "done" else 9,
                               b["pick"]["updated"] if b["pick"]["status"] == "done" else b["pick"]["added"]),
                reverse=False)
    return jsonify(pool=pool, picks=chosen, turn=whose_turn(db, key), season=current_season(), me=user,
                   group=key, members=group_members(key))


@app.route("/api/reto/pick", methods=["POST"])
@login_required
def api_reto_pick():
    """{book_id, status: next|reading|done|null}. null = quitar del reto."""
    data = request.get_json(force=True, silent=True) or {}
    status = data.get("status", "next")
    if status is not None and status not in PICK_STATUS:
        abort(400)
    db = get_db()
    user = session["user"]
    book_id = data.get("book_id")
    key = check_group(user, data.get("group"))
    sql, params = group_pool_sql(key)
    is_match = db.execute(sql + " AND b.id = ?", [*params, book_id]).fetchone()
    other_grp = db.execute("SELECT grp FROM picks WHERE book_id = ? AND grp != ?", (book_id, key)).fetchone()
    if other_grp:
        abort(409)  # ya está en el reto de otro grupo
    if not is_match:
        # Solo libros que queréis leer los dos… salvo para SEARCH_USERS, que pueden apuntar
        # cualquier libro (p. ej. una lectura del reto que ya hicisteis antes de la app).
        if user not in SEARCH_USERS or not db.execute("SELECT 1 FROM books WHERE id = ?", (book_id,)).fetchone():
            abort(404)
    if status is None:
        db.execute("DELETE FROM picks WHERE book_id = ? AND grp = ?", (book_id, key))
    else:
        # Si se apunta directamente como leída (lectura pasada), no gasta turno
        db.execute("INSERT OR IGNORE INTO picks (book_id, picked_by, status, turn, grp) VALUES (?,?,?,?,?)",
                   (book_id, user, status, 0 if status == "done" else 1, key))
        db.execute("UPDATE picks SET status = ?, updated = datetime('now') WHERE book_id = ?", (status, book_id))
    db.commit()
    return jsonify(ok=True, turn=whose_turn(db, key))


# --------------------------------------------------------------------------- #
# Exportar a Excel (solo SEARCH_USERS)
# --------------------------------------------------------------------------- #
DECISION_ES = {"want": "Leer", "reject": "Descartado", "skip": "Pasado", None: ""}
STATUS_ES = {"next": "Próxima", "reading": "Leyendo", "done": "Terminada"}


@app.route("/exportar")
@login_required
def export_xlsx():
    if session["user"] not in SEARCH_USERS:
        abort(403)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    db = get_db()
    users = sorted(USERS)
    rows = db.execute("SELECT * FROM books ORDER BY level, title").fetchall()
    votes = votes_for(db, [r["id"] for r in rows])
    picks = {r["book_id"]: r for r in db.execute("SELECT * FROM picks")}

    def who_read(v):
        readers = [u.capitalize() for u in users if v.get(u, {}).get("read")]
        everyone = "Los dos" if len(users) == 2 else "Todos"
        if len(readers) in (0, len(users)):
            return {0: "Nadie", len(users): everyone}[len(readers)]
        return ", ".join(readers[:-1]) + " y " + readers[-1] if len(readers) > 1 else readers[0]

    def base(r):
        extra = json.loads(r["extra"] or "{}")
        return [_clean(r["ext_id"]), r["title"], r["original"] if r["original"] != r["title"] else "", r["author"],
                r["level"], extra.get("Estación", ""), extra.get("Año", ""), extra.get("Género", ""), r["blurb"]]

    head = ["Nº", "Título", "Título original", "Autor", "Nivel", "Estación", "Año", "Género", "Sinopsis"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Nos lo quedamos"
    ws.append(head + ["Quién lo ha leído", "En el reto"])
    matches = [r for r in rows if all(votes.get(r["id"], {}).get(u, {}).get("decision") == "want" for u in users)]
    season_order = {"Primavera": 0, "Verano": 1, "Otoño": 2, "Invierno": 3}
    matches.sort(key=lambda r: (season_order.get(json.loads(r["extra"] or "{}").get("Estación", ""), 9), r["level"], r["title"]))
    for r in matches:
        p = picks.get(r["id"])
        ws.append(base(r) + [who_read(votes.get(r["id"], {})), STATUS_ES[p["status"]] if p and p["grp"] == "all" else ""])

    ws2 = wb.create_sheet("Todos")
    ws2.append(head + [f"{u.capitalize()} {k}" for u in users for k in ("decide", "leído")] + ["Coinciden"])
    for r in rows:
        v = votes.get(r["id"], {})
        per = []
        for u in users:
            per += [DECISION_ES[v.get(u, {}).get("decision")], "Sí" if v.get(u, {}).get("read") else ""]
        ds = {v.get(u, {}).get("decision") for u in users}
        everyone = "Los dos" if len(users) == 2 else "Todos"
        wanting = [u.capitalize() for u in users if v.get(u, {}).get("decision") == "want"]
        if ds == {"want"}:
            same = f"💞 {everyone} quieren"
        elif ds == {"reject"}:
            same = f"🗑️ {everyone} lo descartan"
        elif len(wanting) >= 2:
            same = "💕 " + ", ".join(wanting[:-1]) + " y " + wanting[-1]
        else:
            same = ""
        ws2.append(base(r) + per + [same])

    # Hoja con todas las lecturas del reto, de todos los grupos
    ws3 = wb.create_sheet("Reto")
    ws3.append(head + ["Grupo", "Estado", "Elegido por"])
    by_id = {r["id"]: r for r in rows}
    for p in sorted(picks.values(), key=lambda p: (p["grp"] if "grp" in p.keys() else "all", p["added"])):
        r = by_id.get(p["book_id"])
        if not r:
            continue
        grp = p["grp"] if "grp" in p.keys() else "all"
        label = ("Los dos" if len(users) == 2 else "Todos") if grp == "all" else \
            " y ".join(m.capitalize() for m in grp.split("+"))
        ws3.append(base(r) + [label, STATUS_ES[p["status"]], p["picked_by"].capitalize()])

    for sheet in (ws, ws2, ws3):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="8A5A3C")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        widths = {"A": 6, "B": 34, "C": 30, "D": 26, "E": 20, "F": 11, "G": 8, "H": 22, "I": 60}
        for col, w in widths.items():
            sheet.column_dimensions[col].width = w
        for col in "JKLMNO":
            sheet.column_dimensions[col].width = 16
        for row in sheet.iter_rows(min_row=2):
            row[8].alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"coffee_and_chapters_{date.today():%Y-%m-%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/health")
def health():
    return "ok"


init_db()

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
